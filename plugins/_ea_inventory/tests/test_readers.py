import csv
import dataclasses
import hashlib
import io
import json

import fitz
import pytest
from openpyxl import Workbook

from plugins._ea_inventory.helpers import readers


def _xlsx(sheets: dict[str, list[list]], merge: dict[str, str] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        if merge and title in merge:
            ws.merge_cells(merge[title])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(*pages: str) -> bytes:
    doc = fitz.open()
    for t in pages:
        page = doc.new_page()
        if t:
            page.insert_text((72, 72), t)
    return doc.tobytes()


def _docx(paragraphs: list[str], table: list[list[str]]) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    t = d.add_table(rows=len(table), cols=len(table[0]))
    for i, row in enumerate(table):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _leaks(doc: readers.SourceDoc, *secrets: str):
    blob = json.dumps(dataclasses.asdict(doc))
    for s in secrets:
        assert s not in blob, f"{s!r} leaked"


# ---- tabular ----------------------------------------------------------------

def test_csv_header_normalised_rows_redacted_empty_rows_skipped():
    data = (
        "Host Name,Mgmt IP,Serial #,Password,Notes,Mgmt IP\n"
        "sw1,10.0.0.1,FOC123,hunter2,snmp-server community Pr1vate RW,x\n"
        ",,,,,\n"
        "sw2,10.0.0.2,FOC456,s3cret,ok,y\n"
    ).encode()
    doc = readers.parse_bytes("inv.csv", data)
    assert doc.kind == "csv" and doc.sheets == [] and not doc.truncated
    assert doc.headers == ["host_name", "mgmt_ip", "serial", "password", "notes", "mgmt_ip_2"]
    assert len(doc.rows) == 2
    assert doc.rows[0]["host_name"] == "sw1" and doc.rows[0]["serial"] == "FOC123"
    assert doc.rows[0]["password"] == "[REDACTED:secret]"
    assert doc.rows[0]["notes"] == "snmp-server community [REDACTED:community] RW"
    assert doc.redactions == {"secret": 2, "community": 1}
    assert doc.source_id == hashlib.sha256(data).hexdigest()
    assert "1\tsw1\t10.0.0.1\tFOC123" in doc.text and "FOC456" in doc.text
    _leaks(doc, "hunter2", "s3cret", "Pr1vate")


@pytest.mark.parametrize("name,sep", [("a.tsv", "\t"), ("a.csv", ";")])
def test_csv_dialects(name, sep):
    doc = readers.parse_bytes(name, f"name{sep}ip\nsw1{sep}10.0.0.1\n".encode())
    assert doc.rows == [{"name": "sw1", "ip": "10.0.0.1"}]


def test_xlsx_merged_cells_empty_rows_numbers_and_sheets():
    data = _xlsx(
        {
            "Switches": [
                [],
                ["Name", "Site", "IP", "SNMP Community", "Serial"],
                ["sw1", "Jakarta", "10.0.0.1", "Pr1vateRW", 12345.0],
                ["sw2", None, "10.0.0.2", "other", 67890],
                [None, None, None, None, None],
                ["sw3", None, "10.0.0.3", "x", "FOC1"],
            ],
            "Firewalls": [["Name", "IP"], ["fw1", "10.9.9.9"]],
        },
        merge={"Switches": "B3:B6"},
    )
    doc = readers.parse_bytes("inv.xlsx", data)
    assert doc.kind == "xlsx" and doc.sheets == ["Switches", "Firewalls"]
    assert doc.headers == ["name", "site", "ip", "snmp_community", "serial"]
    assert [r["name"] for r in doc.rows] == ["sw1", "sw2", "sw3"]
    assert doc.rows[0]["site"] == "Jakarta" and doc.rows[1]["site"] == ""  # merged: tolerated, not filled
    assert doc.rows[0]["serial"] == "12345"  # 12345.0 is not a float to a person
    assert doc.rows[0]["snmp_community"] == "[REDACTED:secret]"
    # the evidence text covers every sheet, the rows only the chosen one
    assert "[sheet Firewalls]" in doc.text and "fw1" in doc.text
    other = readers.parse_bytes("inv.xlsx", data, sheet="Firewalls")
    assert [r["name"] for r in other.rows] == ["fw1"]
    with pytest.raises(ValueError, match="No sheet named 'Nope'.*Switches"):
        readers.parse_bytes("inv.xlsx", data, sheet="Nope")
    _leaks(doc, "Pr1vateRW")


def test_xlsx_first_sheet_without_data_is_skipped():
    doc = readers.parse_bytes("a.xlsx", _xlsx({"Empty": [], "Data": [["Name"], ["sw1"]]}))
    assert doc.rows == [{"name": "sw1"}]


def test_row_cap_truncates(monkeypatch):
    monkeypatch.setattr(readers, "MAX_ROWS", 3)
    doc = readers.parse_bytes("a.csv", ("name\n" + "\n".join(f"sw{i}" for i in range(10)) + "\n").encode())
    assert doc.truncated and len(doc.rows) == 3 and "sw3" not in doc.text


def test_size_cap(monkeypatch, tmp_path):
    monkeypatch.setattr(readers, "MAX_BYTES", 10)
    with pytest.raises(ValueError, match="larger than"):
        readers.parse_bytes("a.txt", b"x" * 11)
    (tmp_path / "big.txt").write_bytes(b"x" * 11)
    with pytest.raises(ValueError, match="larger than"):
        readers.read_source("big.txt", root=tmp_path)


def test_corrupt_xls_and_xlsx_are_clear_errors():
    for name in ("a.xls", "a.xlsx", "a.docx", "a.pdf"):
        with pytest.raises(ValueError, match="Could not read"):
            readers.parse_bytes(name, b"not really a spreadsheet")


# ---- pdf, docx, text --------------------------------------------------------

def test_pdf_pages_markers_and_redaction():
    doc = readers.parse_bytes("net.pdf", _pdf("sw1 10.0.0.1 FOC123", "enable secret 5 $1$ab$cdefgh"))
    assert doc.kind == "pdf" and doc.pages == 2 and doc.rows == []
    assert "[page 1]" in doc.text and "[page 2]" in doc.text and "FOC123" in doc.text
    assert "enable secret 5 [REDACTED:password]" in doc.text and doc.redactions["password"] == 1
    _leaks(doc, "$1$ab$cdefgh")


def test_scanned_pdf_has_a_clear_error():
    with pytest.raises(ValueError, match="no text layer"):
        readers.parse_bytes("scan.pdf", _pdf("", ""))


def test_docx_paragraphs_and_tables():
    doc = readers.parse_bytes(
        "plan.docx",
        _docx(["Core switch sw1", "password: hunter2hunter2"], [["Name", "IP"], ["sw1", "10.0.0.1"]]),
    )
    assert doc.kind == "docx"
    assert "Core switch sw1" in doc.text and "sw1\t10.0.0.1" in doc.text
    assert "password: [REDACTED:secret]" in doc.text
    _leaks(doc, "hunter2hunter2")


@pytest.mark.parametrize("name", ["router.cfg", "r.conf", "r.log", "r.txt", "r.md", "r.json", "r.yaml", "noext"])
def test_text_kinds(name):
    doc = readers.parse_bytes(name, b"hostname r1\nusername a secret 9 $9$hash123\n")
    assert doc.kind == "text" and "hostname r1" in doc.text and "$9$hash123" not in doc.text


def test_binary_and_unsupported_refused():
    with pytest.raises(ValueError, match="binary"):
        readers.parse_bytes("a.txt", b"\x00\x01\x02")
    with pytest.raises(ValueError, match="Unsupported file type"):
        readers.parse_bytes("a.exe", b"MZ")


# ---- containment ------------------------------------------------------------

@pytest.fixture
def tree(tmp_path):
    mine, theirs = tmp_path / "users" / "me", tmp_path / "users" / "other"
    (mine / "uploads").mkdir(parents=True)
    theirs.mkdir(parents=True)
    (mine / "uploads" / "inv.csv").write_text("name\nsw1\n")
    (mine / "notes.txt").write_text("hostname sw1\n")
    (theirs / "secret.txt").write_text("not yours\n")
    (tmp_path / "outside.txt").write_text("nope\n")
    return mine, theirs, tmp_path


def test_reads_inside_by_relative_bare_and_absolute_path(tree):
    mine, _, _ = tree
    assert readers.read_source("inv.csv", root=mine).rows == [{"name": "sw1"}]  # bare upload name
    assert readers.read_source("uploads/inv.csv", root=mine).rows == [{"name": "sw1"}]
    assert readers.read_source(str(mine / "notes.txt"), root=mine).kind == "text"
    assert readers.read_source("notes.txt", root=[mine / "nowhere", mine]).kind == "text"


def test_refuses_outside_traversal_other_user_and_symlinks(tree):
    mine, theirs, base = tree
    (mine / "link.txt").symlink_to(base / "outside.txt")
    (mine / "dirlink").symlink_to(theirs, target_is_directory=True)
    for bad in (
        "../../outside.txt",
        str(base / "outside.txt"),
        str(theirs / "secret.txt"),
        "../other/secret.txt",
        "link.txt",
        "dirlink/secret.txt",
        "/etc/passwd",
    ):
        with pytest.raises(ValueError, match="outside your files"):
            readers.read_source(bad, root=mine)
    # a prefix is not a boundary: .../me-evil is not inside .../me
    (base / "users" / "me-evil").mkdir()
    (base / "users" / "me-evil" / "x.txt").write_text("x")
    with pytest.raises(ValueError, match="outside your files"):
        readers.read_source(str(base / "users" / "me-evil" / "x.txt"), root=mine)


def test_missing_and_empty_and_directory(tree):
    mine, _, _ = tree
    for p in ("missing.csv", "uploads"):
        with pytest.raises(ValueError, match="not found"):
            readers.read_source(p, root=mine)
    with pytest.raises(ValueError, match="No file path"):
        readers.read_source("  ", root=mine)


def test_a_sheet_that_claims_a_million_rows_is_read_only_to_the_cap(monkeypatch):
    """A cell at the far corner made every row and column between be built before the cap applied."""
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "ip"])
    ws.append(["sw1", "10.0.0.1"])
    ws["XFD1048576"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    seen = {}
    real = readers._tabular
    monkeypatch.setattr(readers, "_tabular", lambda n, k, s, sheets, pick: seen.update(sheets) or real(n, k, s, sheets, pick))
    doc = readers.parse_bytes("big.xlsx", buf.getvalue())
    (rows,) = seen.values()
    assert len(rows) <= readers.MAX_ROWS + 1 and len(rows[0]) <= readers.MAX_COLS
    assert doc.rows[0]["name"] == "sw1"


def test_an_archive_that_unpacks_too_large_is_refused(monkeypatch):
    monkeypatch.setattr(readers, "MAX_UNZIPPED", 1000)
    with pytest.raises(ValueError, match="unpacks to more than"):
        readers.parse_bytes("a.xlsx", _xlsx({"s": [["name"], ["x" * 5000]]}))
